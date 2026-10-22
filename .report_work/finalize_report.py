from docx import Document
from docx.enum.text import WD_LINE_SPACING
p=r'C:\Users\yagne\OneDrive\Desktop\PROJECTS\Java_mini_project\Astra_AI_One_Week_Technical_Project_Report.docx'
d=Document(p)
def setp(x,s):
 if x.runs:
  x.runs[0].text=s
  for r in x.runs[1:]: r.text=''
 else: x.add_run(s)
# Put the provisional-day note before the table, rather than after it.
note=next(x for x in d.paragraphs if x.text.startswith('Day-to-activity mapping below'))
d.tables[1]._tbl.addprevious(note._p)
for x in d.paragraphs:
 if x.text.startswith('The JavaFX application maps its main screens'):
  s=x.text.replace('The analysis record contains derived counts and explanations, not the submitted text.','The analysis record contains derived counts and explanations, not the submitted text. The mutable char[] working copy is cleared by the input policy where implemented. JavaFX getText() creates an immutable Java String, which cannot be reliably wiped; therefore the application does not claim complete in-memory erasure. The UI fields are cleared after submission, and the password is not persisted or transmitted.')
  setp(x,s)
 if x.text.startswith('Implementation is organized around'):
  s=x.text.replace('JavaFX password-analysis path.','JavaFX password-analysis path. In a graphical desktop session, the configured launch command is mvn javafx:run; the Maven plugin entry point is com.astraai.Main.')
  setp(x,s)
 if 'HoD SIGNATURE' in x.text:
  x.paragraph_format.page_break_before=False
  x.paragraph_format.keep_with_next=False
  x.paragraph_format.keep_together=False
  x.paragraph_format.space_before=0
  x.paragraph_format.space_after=0
  x.paragraph_format.widow_control=False
d.save(p)
