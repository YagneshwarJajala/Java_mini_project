from docx import Document
p=r'C:\Users\yagne\OneDrive\Desktop\PROJECTS\Java_mini_project\Astra_AI_One_Week_Technical_Project_Report.docx'
d=Document(p)
def setp(p,s):
 if p.runs:
  p.runs[0].text=s
  for r in p.runs[1:]: r.text=''
 else: p.add_run(s)
for x in d.paragraphs:
 if x.text.strip().startswith('[State a clear, specific project objective.]'):
  setp(x,'1. Build a JavaFX workflow with PASSWORD_ENTRY, PROCESSING and FINAL_RESULT states.\n2. Analyze character classes and predictable patterns locally, keeping derived counts in the analysis record.\n3. Present guidance for Argon2id, bcrypt, scrypt, PBKDF2 and an underspecified SHA-3-based entry.\n4. Keep synthetic-index simulation separate from the submitted password.\n5. Display internal educational assumptions and their limitations.\n6. Verify behaviors with tests and Maven builds.')
 if x.text.startswith('Result / Observation :'):
  setp(x,'Result / Observation: Source and automated tests support the three-state flow, local derived-characteristic analysis, separate synthetic simulation, five algorithm cards, four illustrative estimate profiles and optional benchmark/report components. No submitted-password cracking, external network activity, manual JavaFX screenshot or performance-benchmark observation is claimed.')
 if x.text.startswith('The JavaFX application maps its main screens') and '\n\nFigure 1.' in x.text:
  setp(x,x.text.rsplit('\n\nFigure 1.',1)[0])
d.save(p)
